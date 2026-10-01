"""Render a tailored resume to HTML, PDF and DOCX.

All three come from the same structured resume, in a single-column layout with
real text and standard section headings, which is what resume parsers read
most reliably.
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from jinja2 import Environment, FileSystemLoader, select_autoescape
from playwright.sync_api import Browser

from jobportal.browser import launch_browser
from jobportal.resume.model import TailoredResume
from jobportal.settings import Settings
from jobportal.text import slugify

TEMPLATES = Path(__file__).parent / "templates"


@lru_cache
def _environment() -> Environment:
    return Environment(
        loader=FileSystemLoader(TEMPLATES),
        autoescape=select_autoescape(["html", "j2"], default=True),
        trim_blocks=True,
        lstrip_blocks=True,
    )


def render_html(resume: TailoredResume) -> str:
    return _environment().get_template("resume.html.j2").render(resume=resume.model_dump())


def write_pdf(
    html: str, path: Path, *, settings: Settings | None = None, browser: Browser | None = None
) -> Path:
    """Print the HTML to a Letter-size PDF with Chromium."""
    path.parent.mkdir(parents=True, exist_ok=True)

    def _print(active: Browser) -> None:
        page = active.new_page()
        try:
            page.set_content(html, wait_until="load")
            page.pdf(path=str(path), format="Letter", prefer_css_page_size=True)
        finally:
            page.close()

    if browser is not None:
        _print(browser)
    else:
        with launch_browser(settings, headless=True) as owned:
            _print(owned)
    return path


def write_docx(resume: TailoredResume, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.6)
    usable_width = (section.page_width or Inches(8.5)) - Inches(0.7) - Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)

    name = document.add_paragraph()
    run = name.add_run(resume.name)
    run.bold = True
    run.font.size = Pt(20)
    name.paragraph_format.space_after = Pt(0)

    if resume.headline:
        headline = document.add_paragraph()
        headline.add_run(resume.headline).font.size = Pt(11.5)
        headline.paragraph_format.space_after = Pt(0)
    if resume.contact:
        contact = document.add_paragraph("  |  ".join(resume.contact))
        contact.runs[0].font.size = Pt(9.5)

    def heading(text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def bullet(text: str) -> None:
        paragraph = document.add_paragraph(text, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(1.5)

    if resume.summary:
        heading("Summary")
        document.add_paragraph(resume.summary)

    if resume.skills:
        heading("Skills")
        for group in resume.skills:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{group.group}: ").bold = True
            paragraph.add_run(", ".join(group.items))
            paragraph.paragraph_format.space_after = Pt(1.5)

    heading("Experience")
    for role in resume.experience:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(f"{role.title}, {role.company}").bold = True
        paragraph.add_run(f"\t{role.period}")
        if role.location:
            location = document.add_paragraph(role.location)
            location.paragraph_format.space_after = Pt(1)
            location.paragraph_format.keep_with_next = True
        for text in role.bullets:
            bullet(text)

    if resume.education:
        heading("Education")
        for item in resume.education:
            line = ", ".join(part for part in (item.degree, item.school) if part)
            bullet(f"{line} ({item.year})" if item.year else line)

    if resume.certifications:
        heading("Certifications")
        for cert in resume.certifications:
            line = ", ".join(part for part in (cert.name, cert.issuer) if part)
            bullet(f"{line} ({cert.year})" if cert.year else line)

    for extra in resume.extras:
        heading(extra.title)
        for text in extra.items:
            bullet(text)

    document.core_properties.author = resume.name
    document.core_properties.title = f"{resume.name} - Resume"
    document.save(str(path))
    return path


def resume_basename(name: str) -> str:
    """'Alex Example' -> 'Alex_Example_Resume'. The name a recruiter will see."""
    return slugify(name).replace("-", "_").title() + "_Resume"
